"""Template-based XLSX generator for JAMEL estimates ("Wycena").

Renders estimates in the two JAMEL formats:

* ``range`` — widełki min/max: ``Etap | Opis | Zadanie Wykonawcy | Rola |
  RH (min) | RH (max) | Stawka | Kwota całkowita netto (min) | (max)``
  with formulas ``=F*H`` / ``=G*H`` (cols B-J, margins A/K).
* ``fixed`` — cena stała: ``Etap | Opis | Zadanie Wykonawcy | Rola | RH |
  Stawka | Kwota całkowita netto`` with ``=F*G`` (cols B-H, margins A/I).

Output has up to two sheets: ``Wycena`` (core scope) and ``Elementy
dodatkowe`` (upsells, same layout; omitted when there are no
``additional_items``). Visual style — fonts, colors, borders, widths, logo —
comes entirely from ``assets/estimate_template.xlsx`` (built by
``tools/build_template.py``), never hardcoded here.

Line item kinds:
* ``hourly``   — RH × rate, rate resolved from the JAMEL rate card.
* ``fixed``    — ryczałt: literal amount(s) in Kwota, ``-`` in RH/Stawka;
                 a negative amount renders a discount (Rabat).
* ``included`` — "w cenie": visible row at zero cost.

CLI:
    python generate_wycena_xlsx.py input.json -o wycena.xlsx
"""
from __future__ import annotations

import argparse
import json
import math
import sys
import zipfile
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "estimate_template.xlsx"
LOGO_ZIP_PATH = "xl/media/image1.png"
LOGO_WIDTH_PX = 74
LOGO_HEIGHT_PX = 119

CORE_SHEET_TITLE = "Wycena"
ADDITIONAL_SHEET_TITLE = "Elementy dodatkowe"

# ---------------------------------------------------------------------------
# JAMEL rate card (PLN netto/h). Copywriter is deliberately absent — copy is
# priced as a ryczałt (kind="fixed"), never hourly.
# ---------------------------------------------------------------------------
ROLE_RATES: dict[str, Decimal] = {
    "Project Manager": Decimal("225"),
    "Web Designer": Decimal("225"),
    "Front-end Developer": Decimal("250"),
    "Back-end Developer": Decimal("250"),
    "Back-end + Front-end Developer": Decimal("250"),
    "Front-end / Back-end Developer": Decimal("250"),
    "Tester": Decimal("180"),
    "Redaktor": Decimal("180"),
    "Programista": Decimal("250"),
}

VALID_KINDS = {"hourly", "fixed", "included"}


@dataclass(frozen=True)
class ModeSpec:
    """Column layout + template block names for one pricing mode."""

    name: str
    num_cols: int
    rate_col: int            # Stawka
    amount_cols: tuple[int, ...]   # Kwota column(s): (min, max) or (single,)
    hours_cols: tuple[int, ...]    # RH column(s): (min, max) or (single,)

    @property
    def suffix(self) -> str:
        return f"-{self.name}"

    def block(self, base: str) -> str:
        return f"{base}{self.suffix}"

    def item_formulas(self, row: int) -> dict[int, str]:
        """Kwota formulas for an hourly item at ``row``, keyed by column."""
        rate = get_column_letter(self.rate_col)
        return {
            amount_col: f"={get_column_letter(hours_col)}{row}*{rate}{row}"
            for hours_col, amount_col in zip(self.hours_cols, self.amount_cols)
        }


MODES = {
    "range": ModeSpec(name="range", num_cols=11, rate_col=8, amount_cols=(9, 10), hours_cols=(6, 7)),
    "fixed": ModeSpec(name="fixed", num_cols=9, rate_col=7, amount_cols=(8,), hours_cols=(6,)),
}


# ===========================================================================
# Data layer
# ===========================================================================
@dataclass
class Role:
    name: str


@dataclass
class LineItem:
    """Resolved estimate line item (both pricing modes).

    In ``fixed`` mode min == max for hours/totals; ``total_min``/``total_max``
    exist for band verification only — the rendered file carries formulas.
    """

    stage: str
    title: str
    description: str
    kind: str = "hourly"
    hours_min: Decimal | None = None
    hours_max: Decimal | None = None
    hourly_rate: Decimal | None = None
    amount_min: Decimal | None = None
    amount_max: Decimal | None = None
    stage_order: int = 0
    item_order: int = 0
    display_number: str = ""
    total_min: Decimal = Decimal("0")
    total_max: Decimal = Decimal("0")
    role: Role | None = None


def _as_decimal(value: Any) -> Decimal:
    return Decimal(str(value))


def _opt_decimal(item: dict, key: str) -> Decimal | None:
    return _as_decimal(item[key]) if item.get(key) is not None else None


def assemble_line_items(payload: dict, key: str = "line_items") -> list[LineItem]:
    """Resolve raw payload items into ordered, rated :class:`LineItem`s.

    Stage order follows the payload's explicit ``stages`` list (core sheet);
    stages found only in items are appended in first-appearance order.
    Stages listed but itemless are kept as empty groups (banner-only).
    Rate resolution: per-item ``hourly_rate`` > payload ``role_rates`` >
    built-in ``ROLE_RATES`` > ``meta.default_hourly_rate``.
    """
    meta = payload.get("meta", {})
    fallback_rate = _as_decimal(meta.get("default_hourly_rate") or 200)

    rate_map: dict[str, Decimal] = dict(ROLE_RATES)
    for role_name, rate in (payload.get("role_rates") or {}).items():
        rate_map[role_name] = _as_decimal(rate)

    stage_items: dict[str, list[dict]] = defaultdict(list)
    if key == "line_items":
        for stage_name in payload.get("stages") or []:
            stage_items[stage_name]  # register empty stage, preserving order
    for item in payload.get(key) or []:
        if isinstance(item, dict):
            stage_items[item.get("stage", "Programowanie")].append(item)

    result: list[LineItem] = []
    for stage_order, (stage_name, items) in enumerate(stage_items.items(), start=1):
        if not items:
            result.append(LineItem(stage=stage_name, title="", description="",
                                   kind="included", stage_order=stage_order, item_order=0))
            continue
        for item_order, item in enumerate(items, start=1):
            kind = item.get("kind", "hourly")
            if kind not in VALID_KINDS:
                raise ValueError(f"Unknown line item kind: {kind!r} ({item.get('title')})")

            hours_min = _opt_decimal(item, "hours_min")
            hours_max = _opt_decimal(item, "hours_max")
            if item.get("hours") is not None:
                hours_min = hours_max = _as_decimal(item["hours"])

            amount_min = _opt_decimal(item, "amount_min")
            amount_max = _opt_decimal(item, "amount_max")
            if item.get("amount") is not None:
                amount_min = amount_max = _as_decimal(item["amount"])

            role_name = str(item.get("role", ""))
            rate: Decimal | None = None
            if kind == "hourly":
                if item.get("hourly_rate") is not None:
                    rate = _as_decimal(item["hourly_rate"])
                else:
                    rate = rate_map.get(role_name)
                    if rate is None:
                        rate = fallback_rate
                        print(f"WARNING: unknown role {role_name!r} in {item.get('title')!r} "
                              f"-> fallback rate {rate}", file=sys.stderr)
                if hours_min is None or hours_max is None:
                    raise ValueError(f"Hourly item without hours: {item.get('title')!r}")
                total_min, total_max = hours_min * rate, hours_max * rate
            elif kind == "fixed":
                if amount_min is None or amount_max is None:
                    raise ValueError(f"Fixed item without amount: {item.get('title')!r}")
                total_min, total_max = amount_min, amount_max
            else:  # included
                total_min = total_max = Decimal("0")

            result.append(
                LineItem(
                    stage=stage_name,
                    title=item.get("title", ""),
                    description=item.get("description", ""),
                    kind=kind,
                    hours_min=hours_min,
                    hours_max=hours_max,
                    hourly_rate=rate,
                    amount_min=amount_min,
                    amount_max=amount_max,
                    stage_order=stage_order,
                    item_order=item_order,
                    display_number=f"{stage_order}.{item_order}",
                    total_min=total_min,
                    total_max=total_max,
                    role=Role(role_name) if role_name else None,
                )
            )
    return result


def sum_totals(line_items: list[LineItem]) -> tuple[Decimal, Decimal]:
    """Recompute SUMA (min, max) — openpyxl never evaluates formulas."""
    return (
        sum((li.total_min for li in line_items), Decimal("0")),
        sum((li.total_max for li in line_items), Decimal("0")),
    )


# ===========================================================================
# Render layer
# ===========================================================================
def load_template() -> openpyxl.Workbook:
    return openpyxl.load_workbook(TEMPLATE_PATH)


def extract_logo_bytes() -> bytes:
    with zipfile.ZipFile(TEMPLATE_PATH) as zf:
        return zf.read(LOGO_ZIP_PATH)


def _apply_column_widths(source_ws: "Worksheet", target_ws: "Worksheet") -> None:
    for letter, dim in source_ws.column_dimensions.items():
        if dim.width:
            target_ws.column_dimensions[letter].width = dim.width


def _auto_row_height(ws: "Worksheet", row: int, description: str, col_d_width: float) -> None:
    """Grow the row so a wrapped description stays visible."""
    if not description:
        return
    chars_per_line = max(10, int(col_d_width * 1.3))
    lines = sum(
        max(1, math.ceil(len(part) / chars_per_line))
        for part in description.split("\n")
    )
    if lines > 1:
        current = ws.row_dimensions[row].height or 15
        ws.row_dimensions[row].height = max(current, lines * 13)


def _replace_variables(value: Any, var_map: dict[str, Any]) -> Any:
    if not isinstance(value, str):
        return value
    for key, replacement in var_map.items():
        value = value.replace(f"{{{{{key}}}}}", str(replacement))
    return value


def _copy_style(src_cell: Any, dst_cell: Any) -> None:
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format


def _copy_row(
    source_ws: "Worksheet",
    source_row: int,
    target_ws: "Worksheet",
    target_row: int,
    num_cols: int,
    var_map: dict[str, Any] | None = None,
) -> None:
    """Copy one row (values, styles, height, merges) substituting variables."""
    if var_map is None:
        var_map = {}

    for col in range(1, num_cols + 1):
        src_cell = source_ws.cell(row=source_row, column=col)
        dst_cell = target_ws.cell(row=target_row, column=col)
        raw = src_cell.value
        if isinstance(raw, str) and "#REF!" in raw:
            dst_cell.value = None
        else:
            dst_cell.value = _replace_variables(raw, var_map)
        _copy_style(src_cell, dst_cell)

    src_height = source_ws.row_dimensions[source_row].height
    if src_height:
        target_ws.row_dimensions[target_row].height = src_height

    row_delta = target_row - source_row
    existing_ranges = {str(r) for r in target_ws.merged_cells.ranges}
    for merged_range in source_ws.merged_cells.ranges:
        if merged_range.min_row <= source_row <= merged_range.max_row:
            new_min_row = merged_range.min_row + row_delta
            new_max_row = merged_range.max_row + row_delta
            min_col_letter = get_column_letter(merged_range.min_col)
            max_col_letter = get_column_letter(merged_range.max_col)
            target_range = f"{min_col_letter}{new_min_row}:{max_col_letter}{new_max_row}"
            if target_range not in existing_ranges:
                target_ws.merge_cells(target_range)
                existing_ranges.add(target_range)


def _group_by_stage(line_items: list[LineItem]) -> list[tuple[int, str, list[LineItem], list[str]]]:
    """Group items by stage and pick a template block variant per item."""
    stage_map: dict[int, tuple[str, list[LineItem]]] = {}
    for item in line_items:
        so = item.stage_order
        if so not in stage_map:
            stage_map[so] = (item.stage, [])
        if item.item_order > 0:  # item_order == 0 marks a banner-only stage
            stage_map[so][1].append(item)

    groups = []
    for stage_order, (stage_name, items) in sorted(stage_map.items()):
        items.sort(key=lambda li: li.item_order)
        n = len(items)
        if n == 1:
            variants = ["line-item-single"]
        else:
            variants = ["line-item-first"] + ["line-item-middle"] * (n - 2) + ["line-item-last"]
        groups.append((stage_order, stage_name, items, variants))
    return groups


def _write_typed(ws: "Worksheet", row: int, col: int, value: Any,
                 style_src: Any, number_format: str | None = None) -> None:
    """Write a typed value/formula and restore the template cell style."""
    cell = ws.cell(row=row, column=col)
    cell.value = value
    _copy_style(style_src, cell)
    if number_format is not None:
        cell.number_format = number_format


def _write_item_numbers(ws: "Worksheet", row: int, item: LineItem,
                        mode: ModeSpec, block: "Worksheet") -> None:
    """Fill the numeric part of an item row (RH / Stawka / Kwota) per kind."""
    hours_vals = (item.hours_min, item.hours_max)
    amount_vals = (item.amount_min, item.amount_max)

    if item.kind == "hourly":
        for hours_col, hours in zip(mode.hours_cols, hours_vals):
            _write_typed(ws, row, hours_col, float(hours), block.cell(row=1, column=hours_col))
        _write_typed(ws, row, mode.rate_col, float(item.hourly_rate), block.cell(row=1, column=mode.rate_col))
        for amount_col, formula in mode.item_formulas(row).items():
            _write_typed(ws, row, amount_col, formula, block.cell(row=1, column=amount_col))
    elif item.kind == "fixed":
        # Ryczałt: '-' markers, literal amount(s). Negative amount = rabat.
        for hours_col in mode.hours_cols:
            _write_typed(ws, row, hours_col, "-", block.cell(row=1, column=hours_col))
        _write_typed(ws, row, mode.rate_col, "-", block.cell(row=1, column=mode.rate_col))
        for amount_col, amount in zip(mode.amount_cols, amount_vals):
            _write_typed(ws, row, amount_col, float(amount), block.cell(row=1, column=amount_col))
    else:  # included — "w cenie"
        if mode.name == "range":
            for hours_col, hours in zip(mode.hours_cols, hours_vals):
                value = float(hours) if hours is not None else 0.0
                _write_typed(ws, row, hours_col, value, block.cell(row=1, column=hours_col))
            _write_typed(ws, row, mode.rate_col, 0.0, block.cell(row=1, column=mode.rate_col))
            for amount_col in mode.amount_cols:
                _write_typed(ws, row, amount_col, 0.0, block.cell(row=1, column=amount_col))
        else:
            _write_typed(ws, row, mode.hours_cols[0], "W cenie", block.cell(row=1, column=mode.hours_cols[0]))
            _write_typed(ws, row, mode.rate_col, None, block.cell(row=1, column=mode.rate_col))
            _write_typed(ws, row, mode.amount_cols[0], None, block.cell(row=1, column=mode.amount_cols[0]))


def _render_sheet(
    ws: "Worksheet",
    template: openpyxl.Workbook,
    line_items: list[LineItem],
    meta: dict,
    mode: ModeSpec,
) -> None:
    """Render one estimate sheet (header, stages, items, SUMA, spacer)."""
    header_ws = template[mode.block("estimate-header")]
    _apply_column_widths(header_ws, ws)

    author = meta.get("author") or {}
    header_var_map = {
        "estimate-title": meta.get("title", ""),
        "client": meta.get("client_name") or meta.get("title", ""),
        "today": meta.get("date") or date.today().strftime("%d.%m.%Y"),
        "user-firstname": author.get("first_name", ""),
        "user-lastname": author.get("last_name", ""),
    }

    current_row = 1
    for src_row in range(1, header_ws.max_row + 1):
        _copy_row(header_ws, src_row, ws, current_row, mode.num_cols, header_var_map)
        current_row += 1

    try:
        img = XLImage(BytesIO(extract_logo_bytes()))
        img.width = LOGO_WIDTH_PX
        img.height = LOGO_HEIGHT_PX
        img.anchor = "B1"
        ws.add_image(img)
    except Exception:
        pass  # logo is non-critical — render without it on any error

    col_d_width = ws.column_dimensions["D"].width or 40
    first_item_row: int | None = None
    last_item_row: int | None = None

    for _stage_order, stage_name, items, variants in _group_by_stage(line_items):
        _copy_row(template[mode.block("etap-header")], 1, ws, current_row,
                  mode.num_cols, var_map={"stage-name": stage_name})
        current_row += 1

        for item, variant in zip(items, variants):
            block = template[mode.block(variant)]
            var_map = {
                "line-item-number": item.display_number,
                "line-item-title": item.title,
                "line-item-description": item.description,
                "role": item.role.name if item.role else "",
            }
            _copy_row(block, 1, ws, current_row, mode.num_cols, var_map)
            _write_item_numbers(ws, current_row, item, mode, block)
            _auto_row_height(ws, current_row, item.description or "", col_d_width)

            if first_item_row is None:
                first_item_row = current_row
            last_item_row = current_row
            current_row += 1

    # SUMA — =SUM over the contiguous item range (banner rows hold no numbers).
    sum_ws = template[mode.block("sum")]
    formula_row_offset = sum_ws.max_row  # formulas live on the block's last row
    for src_row in range(1, sum_ws.max_row + 1):
        _copy_row(sum_ws, src_row, ws, current_row, mode.num_cols)
        current_row += 1
    if first_item_row is not None:
        formula_row = current_row - 1
        for amount_col in mode.amount_cols:
            letter = get_column_letter(amount_col)
            _write_typed(
                ws, formula_row, amount_col,
                f"=SUM({letter}{first_item_row}:{letter}{last_item_row})",
                sum_ws.cell(row=formula_row_offset, column=amount_col),
            )

    _copy_row(template[mode.block("last-row")], 1, ws, current_row, mode.num_cols)
    ws.sheet_view.showGridLines = False


def generate_xlsx(payload: dict) -> bytes:
    """Render a full estimate workbook from a validated payload."""
    mode_name = payload.get("pricing_mode")
    if mode_name not in MODES:
        raise ValueError(f"pricing_mode must be one of {sorted(MODES)}, got {mode_name!r}")
    mode = MODES[mode_name]
    meta = payload.get("meta", {})

    core_items = assemble_line_items(payload, "line_items")
    additional_items = assemble_line_items(payload, "additional_items")

    template = load_template()
    wb = Workbook()

    ws = wb.active
    ws.title = CORE_SHEET_TITLE
    _render_sheet(ws, template, core_items, meta, mode)

    if any(li.item_order > 0 for li in additional_items):
        ws2 = wb.create_sheet(ADDITIONAL_SHEET_TITLE)
        _render_sheet(ws2, template, additional_items, meta, mode)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_from_payload(payload: dict) -> bytes:
    """Backwards-friendly alias used by the CLI and tests."""
    return generate_xlsx(payload)


# ===========================================================================
# CLI
# ===========================================================================
def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate a JAMEL estimate .xlsx from a JSON payload.")
    parser.add_argument("input", help="Path to the input JSON payload.")
    parser.add_argument("-o", "--output", default="wycena.xlsx", help="Output .xlsx path (default: wycena.xlsx).")
    args = parser.parse_args(argv)

    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
    data = generate_from_payload(payload)
    Path(args.output).write_bytes(data)

    core_min, core_max = sum_totals(assemble_line_items(payload, "line_items"))
    if payload.get("pricing_mode") == "range":
        print(f"Wrote {args.output} ({len(data)} bytes); SUMA core: {core_min}-{core_max} netto")
    else:
        print(f"Wrote {args.output} ({len(data)} bytes); SUMA core: {core_min} netto")
    return 0


if __name__ == "__main__":
    sys.exit(main())
