"""Rebuild assets/estimate_template.xlsx with JAMEL-format building blocks.

Derives mode-suffixed block sheets from the legacy (ELA) blocks so the
template stays reproducible from code:

  fixed  (cols A-I): Etap | Opis | Zadanie Wykonawcy | Rola | RH | Stawka
                     | Kwota całkowita netto            (=F*G)
  range  (cols A-K): Etap | Opis | Zadanie Wykonawcy | Rola | RH (min)
                     | RH (max) | Stawka | Kwota całkowita netto (min)
                     | Kwota całkowita netto (max)      (=F*H / =G*H)

All fonts/fills/borders come from the legacy blocks (already the JAMEL
palette: navy FF21314D, cream FFFEFBF4, coral FFFF818D, Helvetica Neue).
openpyxl drops embedded media on load->save, so the logo PNG is re-injected
into the saved archive via zipfile (the generator reads it raw from
xl/media/image1.png).

Idempotent: on a rebuilt template the fixed blocks serve as the base.

Usage:
    ../scripts/.venv/bin/python build_template.py
"""
from __future__ import annotations

import zipfile
from copy import copy
from pathlib import Path

import openpyxl

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "estimate_template.xlsx"
LOGO_ZIP_PATH = "xl/media/image1.png"

CURRENCY_FMT = "#,##0\\ [$zł-415]"

# Legacy/base block -> the two derived blocks.
BLOCKS = ["estimate-header", "etap-header", "line-item-first", "line-item-middle",
          "line-item-last", "line-item-single", "sum", "last-row"]

# Column remap for range mode: target col index -> source col index.
# Source (A-I): A margin, B etap, C title, D desc, E role, F RH, G rate, H total, I margin.
# Target (A-K): F duplicated into G (RH max), H duplicated into J (Kwota max).
RANGE_COL_MAP = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6, 7: 6, 8: 7, 9: 8, 10: 8, 11: 9}

FIXED_WIDTHS = {"A": 5.75, "B": 9.25, "C": 18.38, "D": 36.63, "E": 10.88,
                "F": 8.25, "G": 7.5, "H": 12.5, "I": 11.13}
RANGE_WIDTHS = {"A": 5.75, "B": 9.25, "C": 18.38, "D": 41.38, "E": 10.88,
                "F": 8.25, "G": 8.25, "H": 10.38, "I": 12.5, "J": 12.5, "K": 11.13}

TITLE_TEXT = "Wycena {{estimate-title}}\ndla {{client}}"

FIXED_HEADER_LABELS = {"B5": "Etap", "C5": "Opis", "D5": "Zadanie Wykonawcy",
                       "E5": "Rola", "F5": "RH", "G5": "Stawka",
                       "H5": "Kwota całkowita netto"}
RANGE_HEADER_LABELS = {"B5": "Etap", "C5": "Opis", "D5": "Zadanie Wykonawcy",
                       "E5": "Rola", "F5": "RH (min)", "G5": "RH (max)",
                       "H5": "Stawka",
                       "I5": "Kwota całkowita netto (min)",
                       "J5": "Kwota całkowita netto (max)"}

# Numeric/amount columns right-aligned like the reference estimates
# (sprint-est/ina-management-est: B,F..J right+top; C/D/E left+top+wrap).
RIGHT_ALIGNED = {"fixed": {2, 6, 7, 8}, "range": {2, 6, 7, 8, 9, 10}}


def _copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format


def _clone_sheet(wb, base_ws, name: str, col_map: dict[int, int], widths: dict[str, float]):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    for row in range(1, base_ws.max_row + 1):
        for tgt_col, src_col in col_map.items():
            _copy_cell(base_ws.cell(row=row, column=src_col), ws.cell(row=row, column=tgt_col))
        height = base_ws.row_dimensions[row].height
        if height:
            ws.row_dimensions[row].height = height
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    return ws


def _right_align(ws, mode: str, rows: range) -> None:
    for row in rows:
        for col in RIGHT_ALIGNED[mode]:
            cell = ws.cell(row=row, column=col)
            alignment = copy(cell.alignment)
            alignment.horizontal = "right"
            cell.alignment = alignment


def _base(wb, block: str):
    """Legacy unsuffixed block if present, else the fixed derivative (re-runs)."""
    return wb[block] if block in wb.sheetnames else wb[f"{block}-fixed"]


def build() -> None:
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    with zipfile.ZipFile(TEMPLATE_PATH) as zf:
        logo_bytes = zf.read(LOGO_ZIP_PATH)

    fixed_map = {i: i for i in range(1, 10)}

    # --- estimate-header ---------------------------------------------------
    base = _base(wb, "estimate-header")
    for mode, col_map, widths, labels, title_merge_end in [
        ("fixed", fixed_map, FIXED_WIDTHS, FIXED_HEADER_LABELS, "H"),
        ("range", RANGE_COL_MAP, RANGE_WIDTHS, RANGE_HEADER_LABELS, "J"),
    ]:
        ws = _clone_sheet(wb, base, f"estimate-header-{mode}", col_map, widths)
        ws["C1"] = TITLE_TEXT
        title_alignment = copy(ws["C1"].alignment)
        title_alignment.wrap_text = True
        ws["C1"].alignment = title_alignment
        for coord, label in labels.items():
            ws[coord] = label
        for row, start in [(1, "C"), (2, "C"), (3, "C"), (4, "B")]:
            ws.merge_cells(f"{start}{row}:{title_merge_end}{row}")
        _right_align(ws, mode, rows=range(5, 6))

    # --- etap-header (coral stage banner) ----------------------------------
    base = _base(wb, "etap-header")
    for mode, col_map, widths, merge_end in [
        ("fixed", fixed_map, FIXED_WIDTHS, "H"),
        ("range", RANGE_COL_MAP, RANGE_WIDTHS, "J"),
    ]:
        ws = _clone_sheet(wb, base, f"etap-header-{mode}", col_map, widths)
        ws.merge_cells(f"B1:{merge_end}1")

    # --- line-item variants -------------------------------------------------
    for variant in ["first", "middle", "last", "single"]:
        base = _base(wb, f"line-item-{variant}")
        # fixed: same layout; normalize Etap format to text and repair the
        # broken '#VALUE!' formula the legacy 'single' block carries.
        ws = _clone_sheet(wb, base, f"line-item-{variant}-fixed", fixed_map, FIXED_WIDTHS)
        ws["B1"].number_format = "General"
        ws["F1"].number_format = "0"
        ws["H1"] = "=F1*G1"
        _right_align(ws, "fixed", rows=range(1, 2))

        # range: F duplicated to G (RH max), H duplicated to J (Kwota max).
        ws = _clone_sheet(wb, base, f"line-item-{variant}-range", RANGE_COL_MAP, RANGE_WIDTHS)
        ws["B1"].number_format = "General"
        ws["F1"] = "{{hours-min}}"
        ws["G1"] = "{{hours-max}}"
        ws["F1"].number_format = "0"
        ws["G1"].number_format = "0"
        ws["I1"] = "=F1*H1"
        ws["J1"] = "=G1*H1"
        for coord in ("I1", "J1"):
            ws[coord].number_format = CURRENCY_FMT
        # I inherited H's right-edge border; the table edge is J now.
        ws["I1"].border = copy(ws["F1"].border)
        _right_align(ws, "range", rows=range(1, 2))

    # --- SUMA ---------------------------------------------------------------
    base = _base(wb, "sum")
    label_style_cell = base["G1"]   # coral bold text on cream
    value_style_cell = base["H1"]   # coral fill, cream bold text
    margin_style_cell = base["A1"]  # plain cream

    ws = _clone_sheet(wb, base, "sum-fixed", fixed_map, FIXED_WIDTHS)
    _right_align(ws, "fixed", rows=range(1, 2))

    # sum-range: 2 rows like sprint-est — labels above coral value cells.
    if "sum-range" in wb.sheetnames:
        del wb["sum-range"]
    ws = wb.create_sheet("sum-range")
    for row in (1, 2):
        for col in range(1, 12):
            _copy_cell(margin_style_cell, ws.cell(row=row, column=col))
            ws.cell(row=row, column=col).value = None
        ws.row_dimensions[row].height = 24.75
    for coord, text in [("I1", "SUMA (min)"), ("J1", "SUMA (max)")]:
        _copy_cell(label_style_cell, ws[coord])
        ws[coord] = text
    for coord in ("I2", "J2"):
        _copy_cell(value_style_cell, ws[coord])
        ws[coord] = None
        ws[coord].number_format = CURRENCY_FMT
    _right_align(ws, "range", rows=range(1, 3))
    for letter, width in RANGE_WIDTHS.items():
        ws.column_dimensions[letter].width = width

    # --- last-row (closing cream spacer) ------------------------------------
    base = _base(wb, "last-row")
    _clone_sheet(wb, base, "last-row-fixed", fixed_map, FIXED_WIDTHS)
    _clone_sheet(wb, base, "last-row-range", RANGE_COL_MAP, RANGE_WIDTHS)

    # --- drop legacy unsuffixed blocks --------------------------------------
    for block in BLOCKS:
        if block in wb.sheetnames:
            del wb[block]

    wb.save(TEMPLATE_PATH)

    # openpyxl drops unreferenced media on save — re-inject the logo so the
    # generator can keep reading it raw from the archive.
    with zipfile.ZipFile(TEMPLATE_PATH, "a", compression=zipfile.ZIP_DEFLATED) as zf:
        if LOGO_ZIP_PATH not in zf.namelist():
            zf.writestr(LOGO_ZIP_PATH, logo_bytes)

    with zipfile.ZipFile(TEMPLATE_PATH) as zf:
        assert LOGO_ZIP_PATH in zf.namelist(), "logo missing after rebuild"
    print(f"Rebuilt {TEMPLATE_PATH.name}: {', '.join(openpyxl.load_workbook(TEMPLATE_PATH).sheetnames)}")


if __name__ == "__main__":
    build()
