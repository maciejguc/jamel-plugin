"""Tests for the standalone wycena xlsx generator.

Two layers of verification:

1. **Golden render fidelity** — ELA's original ``xlsx_generator.py`` has no
   Django imports at module level, so we import it directly, point its
   ``TEMPLATE_PATH`` at the plugin's (byte-identical) bundled template, and feed
   it the SAME resolved line items as our port. The two workbooks must match
   cell-by-cell (values + formulas). This proves the port's render layer stays
   pixel-identical to ELA.

2. **Data layer** — ``assemble_line_items`` must reproduce ELA's
   ``_create_line_items`` logic: stage ordering by insertion, per-stage item
   order, ``display_number = "{stage_order}.{item_order}"``, rate resolution
   (per-item override > role_rates > ROLE_RATES > default), and totals.
"""
from __future__ import annotations

import importlib.util
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

SCRIPTS_DIR = Path(__file__).resolve().parent.parent
PLUGIN_DIR = SCRIPTS_DIR.parent
# Repo root (marketplace dir) is the parent of the jamel-wycena plugin dir; ELA lives beside the repo.
REPO_ROOT = PLUGIN_DIR.parent
TEMPLATE_PATH = PLUGIN_DIR / "assets" / "estimate_template.xlsx"
ELA_GENERATOR = (
    REPO_ROOT.parent
    / "ela" / "backend" / "apps" / "estimates" / "xlsx_generator.py"
)

import generate_wycena_xlsx as gen  # noqa: E402  (port under test)


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------
@pytest.fixture
def payload() -> dict:
    """A deterministic estimate payload covering single + multi-item stages."""
    return {
        "meta": {
            "title": "Strona WWW dla Klienta X",
            "client_name": "Klient X",
            "currency": "PLN",
            "author": {"first_name": "Jan", "last_name": "Kowalski"},
            "default_hourly_rate": 200,
        },
        "profile": "it",
        "role_rates": {"Web Designer": 225},
        "line_items": [
            # analysis: two items (first/last)
            {"stage": "analysis", "title": "Warsztaty UX", "description": "Warsztaty odkrywcze z klientem.",
             "role": "Web Designer", "hours": 16},
            {"stage": "analysis", "title": "Projekt UI", "description": "Projekt graficzny podstron.",
             "role": "Web Designer", "hours": 24},
            # development: one item (single), with per-item rate override
            {"stage": "development", "title": "Frontend - strona główna", "description": "Kodowanie strony głównej.",
             "role": "Front-end Developer", "hours": 20, "hourly_rate": 300},
            # delivery: one item (single), role uses ROLE_RATES fallback
            {"stage": "delivery", "title": "Testy", "description": "Testy funkcjonalne.",
             "role": "Tester", "hours": 8},
        ],
    }


def _load_ela_generator():
    """Import ELA's xlsx_generator by path and repoint it at the plugin template."""
    spec = importlib.util.spec_from_file_location("ela_xlsx_generator", ELA_GENERATOR)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.TEMPLATE_PATH = TEMPLATE_PATH
    return mod


class _FakeManager:
    def __init__(self, items):
        self._items = items

    def select_related(self, *_a, **_k):
        return self

    def order_by(self, *_a, **_k):
        return list(self._items)


class _FakeUser:
    def __init__(self, full_name):
        self._full_name = full_name

    def get_full_name(self):
        return self._full_name


class _FakeEstimate:
    def __init__(self, title, client_name, full_name, items):
        self.title = title
        self.client_name = client_name
        self.created_by = _FakeUser(full_name)
        self.line_items = _FakeManager(items)


# --------------------------------------------------------------------------
# Data layer
# --------------------------------------------------------------------------
def test_assemble_orders_and_numbers(payload):
    items, _meta = gen.assemble_line_items(payload)

    assert [li.display_number for li in items] == ["1.1", "1.2", "2.1", "3.1"]
    assert [li.stage for li in items] == ["analysis", "analysis", "development", "delivery"]
    assert [(li.stage_order, li.item_order) for li in items] == [(1, 1), (1, 2), (2, 1), (3, 1)]


def test_assemble_resolves_rates(payload):
    items, _meta = gen.assemble_line_items(payload)
    by_title = {li.title: li for li in items}

    # role_rates override
    assert by_title["Warsztaty UX"].hourly_rate == Decimal("225")
    # per-item hourly_rate override wins over everything
    assert by_title["Frontend - strona główna"].hourly_rate == Decimal("300")
    # ROLE_RATES fallback (Tester => 180)
    assert by_title["Testy"].hourly_rate == Decimal("180")


def test_assemble_computes_totals(payload):
    items, _meta = gen.assemble_line_items(payload)
    by_title = {li.title: li for li in items}

    assert by_title["Warsztaty UX"].total == Decimal("16") * Decimal("225")
    assert by_title["Frontend - strona główna"].total == Decimal("20") * Decimal("300")


def test_default_rate_fallback():
    """A role absent from role_rates and ROLE_RATES falls back to default_hourly_rate."""
    payload = {
        "meta": {"title": "T", "client_name": "C", "default_hourly_rate": 200,
                 "author": {"first_name": "A", "last_name": "B"}},
        "line_items": [
            {"stage": "development", "title": "Egzotyczna rola", "description": "x",
             "role": "Astronauta", "hours": 10},
        ],
    }
    items, _meta = gen.assemble_line_items(payload)
    assert items[0].hourly_rate == Decimal("200")


# --------------------------------------------------------------------------
# Golden render fidelity
# --------------------------------------------------------------------------
def test_sheet_title_and_dimensions(payload):
    data = gen.generate_from_payload(payload)
    wb = openpyxl.load_workbook(BytesIO(data))
    ws = wb["Wycena"]

    # header(5) + analysis(etap + 2) + development(etap + 1) + delivery(etap + 1) + sum(1) + last-row(1)
    expected_rows = 5 + (1 + 2) + (1 + 1) + (1 + 1) + 1 + 1
    assert ws.max_row == expected_rows


def test_header_variables_substituted(payload):
    data = gen.generate_from_payload(payload)
    text = " ".join(
        str(c.value)
        for row in openpyxl.load_workbook(BytesIO(data))["Wycena"].iter_rows()
        for c in row
        if c.value is not None
    )
    assert "Strona WWW dla Klienta X" in text
    assert "Jan" in text and "Kowalski" in text
    assert "{{" not in text  # no unsubstituted placeholders remain


def test_item_and_sum_formulas(payload):
    data = gen.generate_from_payload(payload)
    ws = openpyxl.load_workbook(BytesIO(data))["Wycena"]

    item_formulas = [
        c.value for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=F")
    ]
    assert len(item_formulas) == 4  # one =Fn*Gn per line item
    for f in item_formulas:
        assert f.count("*") == 1 and f.startswith("=F") and "*G" in f

    sum_formulas = [
        c.value for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=SUM(")
    ]
    assert len(sum_formulas) == 1
    assert sum_formulas[0].count("H") == 4  # all four item H cells summed


def test_matches_ela_golden_render(payload):
    """Same resolved line items through both generators → identical cells."""
    ela = _load_ela_generator()

    items, meta = gen.assemble_line_items(payload)
    mine = gen.generate_from_payload(payload)

    estimate = _FakeEstimate(
        title=meta["title"],
        client_name=meta["client_name"],
        full_name=f"{meta['author']['first_name']} {meta['author']['last_name']}",
        items=items,
    )
    golden = ela.generate_xlsx(estimate)

    ws_mine = openpyxl.load_workbook(BytesIO(mine))["Wycena"]
    ws_gold = openpyxl.load_workbook(BytesIO(golden))["Wycena"]

    assert ws_mine.max_row == ws_gold.max_row
    assert ws_mine.max_column == ws_gold.max_column
    for row in range(1, ws_gold.max_row + 1):
        for col in range(1, ws_gold.max_column + 1):
            assert ws_mine.cell(row=row, column=col).value == ws_gold.cell(row=row, column=col).value, (
                f"cell mismatch at r{row}c{col}"
            )
